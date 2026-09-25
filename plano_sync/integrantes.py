"""Clasificacion por integrante/equipo a partir de 'Equipos e integrantes - VENTAS.xlsx'.

Replica la consulta EquiposIntegrantesVENTAS: hoja 'Integrantes', columnas Equipo e Integrante,
filas validas (ambos no vacios), llave = nombre normalizado (mayusculas, sin acentos), distinto por llave.
Extra: si el nombre en ACC no coincide exacto, se intenta por subconjunto de palabras (solo si es unico).
"""
import re
import unicodedata

SIN_INTEGRANTE = "Sin integrante asignado"
SIN_EQUIPO = "Sin equipo"


def normalizar(v):
    """Igual que NormalizarNombre de Power Query, pero tambien unifica la n con tilde y colapsa espacios."""
    s = re.sub(r"[\x00-\x1f\x7f]", "", str(v or "")).strip().upper()
    s = "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", s)


def cargar_equipos(ruta_xlsx, hoja="Integrantes"):
    from openpyxl import load_workbook
    ws = load_workbook(ruta_xlsx, data_only=True, read_only=True)[hoja]
    filas = list(ws.iter_rows(values_only=True))
    enc = [str(c or "").strip() for c in filas[0]]
    ie, ii = enc.index("Equipo"), enc.index("Integrante")
    out, vistos = [], set()
    for r in filas[1:]:
        eq = re.sub(r"[\x00-\x1f\x7f]", "", str(r[ie] or "")).strip()
        it = re.sub(r"[\x00-\x1f\x7f]", "", str(r[ii] or "")).strip()
        if not eq or not it:
            continue
        k = normalizar(it)
        if k in vistos:
            continue
        vistos.add(k)
        out.append({"equipo": eq, "integrante": it, "key": k, "tokens": set(k.split())})
    return out


class Clasificador:
    def __init__(self, equipos):
        self.equipos = equipos
        self.por_key = {e["key"]: e for e in equipos}

    def asignar(self, nombre_subido):
        """-> (integrante, equipo, tipo_coincidencia)"""
        k = normalizar(nombre_subido)
        if not k:
            return SIN_INTEGRANTE, SIN_EQUIPO, "sin nombre"
        e = self.por_key.get(k)
        if e:
            return e["integrante"], e["equipo"], "exacta"
        toks = set(k.split())
        cand = [x for x in self.equipos if x["tokens"] <= toks or toks <= x["tokens"]]
        if len(cand) == 1 and min(len(cand[0]["tokens"]), len(toks)) >= 2:
            return cand[0]["integrante"], cand[0]["equipo"], "parcial"
        return SIN_INTEGRANTE, SIN_EQUIPO, "sin coincidencia"
