"""Genera index.html (mismo estilo que el reporte 'Planos Forma vs Revit' del portal), data.json y comparativa.xlsx."""
import json
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

from .compare import SIN_MODELO
from .parsing import normalizar_numero
from .integrantes import SIN_INTEGRANTE


def _sn(b):
    return "Sí" if b else "No"


def _fecha(v):
    return str(v)[:10] if v else None


def construir_dashboard(comp, faltantes, avisos, titulo, subtitulo):
    detail = []
    for r in comp:
        detail.append({
            "Proyecto": r["proyecto"], "Numero de plano": r["numero"], "Nombre de plano": r["nombre"],
            "Disciplina": r["disciplina"], "Archivo Forma": r["archivo"], "Formato": r["formato"],
            "Estado comparativa": r["estado"], "Viene de Revit": _sn(r["viene_de_revit"]),
            "Coincidencia exacta": _sn(r["exacta"]), "Numero Revit": r["numero_revit"],
            "Nombre Revit": r["nombre_revit"], "Coincidencias Revit": r["coincidencias"],
            "Integrante": r.get("integrante"), "Equipo": r.get("equipo"), "Modelo Revit": r.get("modelo_revit"),
            "Subido por (PDF)": r.get("subido_por"),
            "Origen PDF": r.get("origen_pdf"), "Ruta Forma": r.get("ruta"), "Última actualización": _fecha(r.get("actualizado")),
        })
    claves = Counter((r["proyecto"], normalizar_numero(r["numero"])) for r in comp if r["numero"])
    duplicadas = sum(1 for r in comp if r["numero"] and claves[(r["proyecto"], normalizar_numero(r["numero"]))] > 1)
    return {
        "generatedAt": datetime.now(timezone.utc).isoformat(), "titulo": titulo, "subtitulo": subtitulo,
        "detail": detail, "duplicateRows": duplicadas,
        "multipleCandidates": sum(1 for r in comp if r["coincidencias"] > 1),
        "sinIntegrante": sum(1 for r in comp if r.get("integrante") == SIN_INTEGRANTE),
        "proyectosSinModelo": sorted({r["proyecto"] for r in comp if r["estado"] == SIN_MODELO}),
        "faltantes": [{"Proyecto": f["proyecto"], "Numero": f["numero"], "Nombre": f["nombre"]} for f in faltantes],
        "avisos": avisos,
    }


def escribir_todo(out_dir, comp, faltantes, res, avisos, titulo, subtitulo="Forma vs Revit"):
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    dash = construir_dashboard(comp, faltantes, avisos, titulo, subtitulo)
    data = json.dumps(dash, ensure_ascii=False, default=str)
    (out / "data.json").write_text(data, encoding="utf-8")
    _xlsx(out / "comparativa.xlsx", comp, faltantes, res)
    html = (Path(__file__).parent / "dashboard.html").read_text(encoding="utf-8")
    (out / "index.html").write_text(
        html.replace("__TITULO__", titulo).replace("__DATA__", data.replace("</", "<\\/")), encoding="utf-8")
    return out


def _xlsx(path, comp, faltantes, res):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    ws.append(["Indicador", "Valor"])
    for k, v in res.items():
        ws.append([k, v])
    cols = [("proyecto", "Proyecto"), ("numero", "Numero de plano"), ("nombre", "Nombre de plano"),
            ("disciplina", "Disciplina"), ("integrante", "Integrante (autor del modelo Revit)"), ("equipo", "Equipo"),
            ("modelo_revit", "Modelo Revit"), ("subido_por", "Subido por el PDF (referencia)"),
            ("archivo", "Archivo Forma"), ("formato", "Formato"),
            ("ruta", "Ruta Forma"), ("actualizado", "Ultima actualizacion"), ("estado", "Estado comparativa"),
            ("numero_revit", "Numero Revit"), ("nombre_revit", "Nombre Revit"), ("viene_de_revit", "Viene de Revit"),
            ("exacta", "Coincidencia exacta"), ("coincidencias", "Coincidencias Revit"), ("origen_pdf", "Origen PDF (metadatos)")]
    w2 = wb.create_sheet("Comparativa")
    w2.append([c[1] for c in cols])
    for r in comp:
        w2.append([r.get(c[0]) for c in cols])
    w3 = wb.create_sheet("Revit sin publicar en Forma")
    w3.append(["Proyecto", "Numero de plano", "Nombre de plano"])
    for r in faltantes:
        w3.append([r["proyecto"], r["numero"], r["nombre"]])
    for w in (ws, w2, w3):
        for c in w[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F4E79")
        w.freeze_panes = "A2"
        for col in w.columns:
            w.column_dimensions[col[0].column_letter].width = min(60, max(12, max(len(str(c.value or "")) for c in col[:200]) + 2))
    w2.auto_filter.ref = w2.dimensions
    wb.save(path)
